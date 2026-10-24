"""Real Word/Excel PDF export plus administrator switching. No physical print jobs."""
import http.cookiejar
import argparse
import io
import json
import os
from pathlib import Path
import secrets
import socket
import subprocess
import sys
import time
import urllib.error
import urllib.request
import uuid
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.test-tools'))
from openpyxl import Workbook
import xlwt
parser = argparse.ArgumentParser()
parser.add_argument('--package', type=Path)
args = parser.parse_args()
TEST = ROOT / '.test-artifacts' / ('office-' + uuid.uuid4().hex)
TEST.mkdir(parents=True)
with socket.socket() as sock:
    sock.bind(('127.0.0.1', 0))
    port = sock.getsockname()[1]
BASE = f'http://127.0.0.1:{port}'
secret = secrets.token_urlsafe(24)
env = dict(os.environ, Storage__Path=str(TEST / 'data 中文'), WEBPRINT_ADMIN_PASSWORD=secret,
           Printing__Enabled='false', Printing__LibreOfficePath=str(TEST / 'not-installed/soffice.exe'),
           Printing__OfficeConverter='libreoffice')
if args.package:
    with zipfile.ZipFile(args.package) as z:
        z.extractall(TEST / 'app')
    command = [str(TEST / 'app/WebPrint.exe')]
    env.update(PATH=os.path.join(os.environ['SystemRoot'], 'System32'), DOTNET_ROOT=str(TEST / 'no-dotnet'))
else:
    command = ['dotnet', str(ROOT / 'bin/Debug/net9.0/WebPrint.dll')]
log = (TEST / 'server-private.log').open('w', encoding='utf-8')
process = None

def client():
    return urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))

def request(opener, path, method='GET', data=None, headers=None, expected=200):
    h = {'X-WebPrint': '1'}
    if isinstance(data, dict):
        h['Content-Type'] = 'application/json'
        data = json.dumps(data).encode()
    h.update(headers or {})
    try:
        r = opener.open(urllib.request.Request(BASE + path, data=data, headers=h, method=method), timeout=115)
    except urllib.error.HTTPError as e:
        r = e
    body = r.read()
    assert r.code == expected, (path, r.code, body.decode(errors='replace'))
    return json.loads(body) if 'application/json' in r.headers.get('Content-Type', '') else body

def start():
    global process
    process = subprocess.Popen([*command, '--urls', BASE],
                               cwd=ROOT, env=env, stdout=log, stderr=log)
    for _ in range(100):
        if process.poll() is not None:
            raise RuntimeError('Server exited; inspect private log locally')
        try:
            urllib.request.urlopen(BASE, timeout=1).close()
            return
        except OSError:
            time.sleep(.2)
    raise RuntimeError('Startup timeout')

def stop():
    if process is not None:
        process.terminate()
        process.wait(timeout=15)

def package(entries):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, content in entries.items():
            z.writestr(name, content)
    return stream.getvalue()

def docx():
    return package({
        '[Content_Types].xml': '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
        '_rels/.rels': '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
        'word/document.xml': '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>WebPrint 中文转换测试 - Page one</w:t></w:r></w:p><w:p><w:r><w:br w:type="page"/></w:r></w:p><w:p><w:r><w:t>WebPrint - Page two</w:t></w:r></w:p><w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr></w:body></w:document>',
    })

def xlsx():
    book = Workbook()
    book.active.title = 'First'
    book.create_sheet('Second')
    for sheet in book:
        sheet['A1'] = 'WebPrint 中文转换测试'
        sheet['Z100'] = 'Outside print area - must not be exported'
        sheet.print_area = 'A1:B3'
    stream = io.BytesIO(); book.save(stream)
    return stream.getvalue()

def xls():
    book = xlwt.Workbook()
    for name in ['First', 'Second']:
        book.add_sheet(name).write(0, 0, 'WebPrint 中文转换测试')
    stream = io.BytesIO(); book.save(stream)
    return stream.getvalue()

def upload(opener, name, content, expected=200):
    boundary = 'OfficeTestBoundary'
    body = f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{name}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode() + content + f'\r\n--{boundary}--\r\n'.encode()
    return request(opener, '/api/documents', 'POST', body, {'Content-Type': f'multipart/form-data; boundary={boundary}'}, expected)

try:
    start()
    anonymous, admin, user = client(), client(), client()
    request(anonymous, '/api/conversion', expected=401)
    request(admin, '/api/login', 'POST', {'username': 'admin', 'password': secret})
    request(admin, '/api/users', 'POST', {'username': 'officetest', 'password': secret, 'role': 'user'})
    request(user, '/api/login', 'POST', {'username': 'officetest', 'password': secret})
    request(user, '/api/conversion', expected=403)
    request(user, '/api/conversion', 'POST', {'provider': 'microsoft-office'}, expected=403)
    request(admin, '/api/conversion', 'POST', {'provider': 'invalid'}, expected=400)
    request(admin, '/api/conversion', 'POST', {'provider': 'microsoft-office'}, {'X-WebPrint': ''}, 400)
    status = request(admin, '/api/conversion')
    assert status['provider'] == 'libreoffice' and not status['libreOffice']
    assert status['microsoftWord'] and status['microsoftExcel'], 'This test requires installed desktop Word and Excel'
    assert 'LibreOffice' in upload(user, 'first.docx', docx(), 503)['error']
    request(admin, '/api/conversion', 'POST', {'provider': 'microsoft-office'})
    for name, content in [('中文文档 with spaces.docx', docx()), ('中文表格 with spaces.xlsx', xlsx()), ('legacy.xls', xls())]:
        document = upload(user, name, content)
        assert document['pages'] == 2, (name, document)
        preview = request(user, f"/api/documents/{document['id']}/pdf")
        assert preview.startswith(b'%PDF')
        print('PASS: actual Microsoft Office export, 2 pages:', name)
    request(admin, '/api/conversion', 'POST', {'provider': 'libreoffice'})
    assert request(user, f"/api/documents/{document['id']}/pdf") == preview, 'Switch changed existing preview'
    assert 'LibreOffice' in upload(user, 'after-switch.docx', docx(), 503)['error']
    request(admin, '/api/conversion', 'POST', {'provider': 'microsoft-office'})
    stop(); start()
    admin, user = client(), client()
    request(admin, '/api/login', 'POST', {'username': 'admin', 'password': secret})
    request(user, '/api/login', 'POST', {'username': 'officetest', 'password': secret})
    assert request(admin, '/api/conversion')['provider'] == 'microsoft-office'
    assert upload(user, 'after-restart.docx', docx())['pages'] == 2
    events = request(admin, '/api/events')
    assert len([e for e in events if e['action'] == 'converter_changed']) == 3
    assert request(user, '/api/jobs') == []
    print('PASS: admin-only switching, validation/CSRF, explicit provider routing, restart persistence, audit, unchanged previews; no print jobs')
    print('Artifacts:', TEST)
finally:
    stop()
    log.close()
